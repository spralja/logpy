from .base_controller import BaseController

from ..model import Entry

from pathlib import Path
from typing import List, Tuple, Optional
from datetime import datetime, timedelta, timezone
import re

import openpyxl
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet


class ExcelController(BaseController):
    DAYS = (None, 'Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday',)
    DAY_RANGE = range(1, 8)

    def __init__(self, log: Path):
        self.entries = ExcelController.load_log(log)
        self.entries.sort()

    @classmethod
    def load_log(cls, log: Path) -> List[Entry]:
        entries = []
        for year in log.iterdir():
            pattern = re.compile(r'^(0|[1-9][0-9]*)$')

            if not pattern.match(year.name):
                #print("skipping-year ", year)
                continue

            if year.is_file():
                raise "ERROR"

            year_number = int(year.name)

            #print("loading-year ", year_number)

            start_time = datetime.fromisocalendar(year_number, 1, 1)

            entries += cls.load_year(year, start_time=start_time)

        return entries

    @classmethod
    def load_year(cls, year: Path, *, start_time: datetime) -> List[Entry]:
        year_number, *_ = start_time.isocalendar()

        entries = []
        for week in year.iterdir():
            if week.is_dir():
                raise "ERROR"

            pattern = re.compile(r'^w\d{2}\.xlsx$')

            if not pattern.match(week.name):
                #print("skipping   ", week)
                continue

            #print("processing ", week)

            workbook = openpyxl.load_workbook(week, read_only=True)

            week_number = int(week.name[1:3])

            start_time = datetime.fromisocalendar(year_number, week_number, 1)

            try:
                entries += cls.load_workbook(workbook, start_time=start_time)
            except ValueError as e:
                print("skipping   ", week, " - invalid format")
                print()
                print(e)
                print()

        return entries

    @classmethod
    def load_workbook(cls, workbook: Workbook, *, start_time: datetime) -> List[Entry]:
        year, week, _ = start_time.isocalendar()

        entries = []
        for day in cls.DAY_RANGE:
            import pytz
            cph_tz = pytz.timezone('Europe/Copenhagen')
            #start_time = cph_tz.localize(datetime.fromisocalendar(year, week, day), is_dst=None)

            start_time = datetime.fromisocalendar(year, week, day)

            sheet_name = cls.DAYS[day]

            sheet = workbook[sheet_name]

            #print("-- loading ", sheet_name)
            try:
                entries += cls.load_sheet(sheet, start_time=start_time)
            except ValueError as e:
                print("---invalid workbook")
                print(e)
                print('---')
                raise ValueError("Invalid workbook")

        return entries

    @classmethod
    def load_sheet(cls, sheet: Worksheet, *, start_time: datetime) -> List[Entry]:
        entries = []
        for row in sheet.iter_rows(min_row=3, values_only=True):
            if row[0] is None:
                break

            start_tod, end_tod, category, description, *optional = row
            meta_tzdiff = None
            if optional: meta_tzdiff, *_ = optional
            tzdiff = timedelta(hours=0)
            pattern = re.compile(r'^__LOGPY_\d{3}$')
            if isinstance(meta_tzdiff, str) and pattern.match(meta_tzdiff):
                #print(meta_tzdiff)
                tzdiff = -timedelta(hours=int(meta_tzdiff[8:11]) / 60)

            sdt_is_dst = None
            edt_is_dst = None
            if meta_tzdiff == '__LOGPY2_DST':
                sdt_is_dst = True
                edt_is_dst = True

            if meta_tzdiff == '__LOGPY2_NDST':
                sdt_is_dst = False
                edt_is_dst = False

            if meta_tzdiff == '__LOGPY2_DST_TO_NDST':
                sdt_is_dst = True
                edt_is_dst = False

            #print(start_tod, end_tod, category, description)
            start_timedelta = timedelta(hours=start_tod.hour, minutes=start_tod.minute)
            end_timedelta = timedelta(hours=end_tod.hour, minutes=end_tod.minute)
            if end_timedelta.total_seconds() == 0:
                end_timedelta = timedelta(days=1)

            start_timedelta += tzdiff
            end_timedelta += tzdiff
            try:
                import pytz
                cph_tz = pytz.timezone('Europe/Copenhagen')
                entry_start_time = cph_tz.localize(start_time + start_timedelta, is_dst=sdt_is_dst)
                entry_end_time = cph_tz.localize(start_time + end_timedelta, is_dst=edt_is_dst)


                entries.append(Entry(
                    entry_start_time,
                    entry_end_time,
                    category,
                    description
                ))
                #print(entries[-1])
                #if tzdiff < timedelta(hours=0): print(entries[-1])
            except ValueError as e:
                #raise "ERRE"
                print("-----invalid entry----------")
                print("------------------------------------------------------------------------------")
                print("------------------------------------------------------------------------------")
                print("------------------------------------------------------------------------------")
                print("------------------------------------------------------------------------------")

                print(e)
                print("------")
                print(start_time + start_timedelta,
                    start_time + end_timedelta,
                    category,
                    description if description else '', tzdiff, sep=', ')
                raise ValueError("Invalid Entry")

        if not entries:
            pass #print(start_time.isocalendar())
        return entries

    def _find_first_after(self, dt: datetime) -> Optional[Entry]:
        for entry in self.entries:
            if entry.start_time > dt:
                return entry

        return None

    def _find_last_before(self, dt: datetime) -> Optional[Entry]:
        for entry in self.entries:
            if entry.end_time > dt:
                return entry

        return None

    def get_intersection(self, start_time, end_time) -> Tuple[Entry]:
        """
        Creates a sorted tuple of entries that intersect with the interval

        :param start_time: The start time of the interval (must be utc)
        :type start_time: datetime.datetime
        :param end_time: The end time of the interval (must be utc)
        :type end_time: datetime.datetime
        :return: A sorted tuple of entry intersections with the interval
        :rtype: tuple[model.Entry]
        """

        # check that start_time is utc
        if start_time.tzinfo != timezone.utc:
            raise ValueError('start_time.tzinfo must be utc')

        # check that end_time is utc
        if end_time.tzinfo != timezone.utc:
            raise ValueError('end_time.tzinfo must be utcc')

        # the list of entries to be returned
        entries = []

        for entry in self.entries:
            intersection = entry.intersection(start_time, end_time)
            if intersection is None:
                continue

            entries.append(intersection)

        return tuple(entries)
